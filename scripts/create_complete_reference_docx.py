import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_complete_reference():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    C_PRIMARY = RGBColor(14, 116, 144)      # Cyan / Teal (#0E7490)
    C_DARK = RGBColor(15, 23, 42)          # Slate-900 (#0F172A)
    C_MUTED = RGBColor(100, 116, 139)      # Slate-500 (#64748B)

    HEX_PRIMARY = "0E7490"
    HEX_DARK = "0F172A"
    HEX_SLATE_50 = "F8FAFC"
    HEX_SLATE_100 = "F1F5F9"
    HEX_SLATE_200 = "E2E8F0"
    HEX_AMBER_LIGHT = "FEF3C7"
    HEX_AMBER_BORDER = "F59E0B"
    HEX_CYAN_LIGHT = "E0F2FE"

    def set_cell_shading(cell, color_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
        tcPr = cell._tc.get_or_add_tcPr()
        def b_tag(side, cfg):
            if not cfg: return f'<w:{side} w:val="none"/>'
            return f'<w:{side} w:val="{cfg.get("val","single")}" w:sz="{cfg.get("sz","4")}" w:space="0" w:color="{cfg.get("color","auto")}"/>'
        xml = f'''<w:tcBorders {nsdecls("w")}>
            {b_tag("top", top)}
            {b_tag("left", left)}
            {b_tag("bottom", bottom)}
            {b_tag("right", right)}
        </w:tcBorders>'''
        tcPr.append(parse_xml(xml))

    def add_callout(text, title="BEGINNER TIP", icon="💡"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(7.3)
        cell = tbl.cell(0, 0)
        set_cell_shading(cell, HEX_AMBER_LIGHT)
        set_cell_margins(cell, top=100, bottom=100, left=180, right=160)
        set_cell_borders(cell, left={"val": "single", "sz": "20", "color": HEX_AMBER_BORDER})
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_icon = p.add_run(f"{icon} {title}: ")
        r_icon.font.name = "Segoe UI"
        r_icon.font.bold = True
        r_icon.font.size = Pt(9.5)
        r_icon.font.color.rgb = C_DARK

        r_text = p.add_run(text)
        r_text.font.name = "Segoe UI"
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = C_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_section_header(title, subtitle=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.font.name = "Segoe UI"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = C_PRIMARY

        if subtitle:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(6)
            r_sub = p_sub.add_run(subtitle)
            r_sub.font.name = "Segoe UI"
            r_sub.font.size = Pt(9.5)
            r_sub.font.color.rgb = C_MUTED

    def add_function_table(rows_data):
        tbl = doc.add_table(rows=len(rows_data) + 1, cols=3)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(1.8)
        tbl.columns[1].width = Inches(2.7)
        tbl.columns[2].width = Inches(2.8)

        # Header Row
        headers = ["Function / Control", "What It Does (Plain English)", "How to Use It"]
        for i, h in enumerate(headers):
            c = tbl.cell(0, i)
            set_cell_shading(c, HEX_DARK)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(h)
            r.font.name = "Segoe UI"
            r.font.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, (fn, what, how) in enumerate(rows_data, start=1):
            c0 = tbl.cell(r_idx, 0)
            c1 = tbl.cell(r_idx, 1)
            c2 = tbl.cell(r_idx, 2)
            bg = HEX_SLATE_50 if r_idx % 2 == 0 else "FFFFFF"
            for c in (c0, c1, c2):
                set_cell_shading(c, bg)
                set_cell_margins(c, top=70, bottom=70, left=100, right=100)
                set_cell_borders(c, bottom={"val": "single", "sz": "4", "color": HEX_SLATE_200})

            # Col 0: Name
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(fn)
            r0.font.name = "Segoe UI"
            r0.font.bold = True
            r0.font.size = Pt(9)
            r0.font.color.rgb = C_PRIMARY

            # Col 1: What it does
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(what)
            r1.font.name = "Segoe UI"
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = C_DARK

            # Col 2: How to use
            p2 = c2.paragraphs[0]
            r2 = p2.add_run(how)
            r2.font.name = "Segoe UI"
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = C_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==============================================================
    # 1. HEADER BANNER
    # ==============================================================
    tbl_hero = doc.add_table(rows=1, cols=1)
    tbl_hero.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_hero.autofit = False
    tbl_hero.columns[0].width = Inches(7.3)
    c_hero = tbl_hero.cell(0, 0)
    set_cell_shading(c_hero, HEX_DARK)
    set_cell_margins(c_hero, top=200, bottom=200, left=240, right=240)
    set_cell_borders(c_hero, bottom={"val": "single", "sz": "20", "color": HEX_PRIMARY})

    p_badge = c_hero.paragraphs[0]
    p_badge.paragraph_format.space_after = Pt(2)
    r_badge = p_badge.add_run("CADENCE DAW  •  COMPLETE FUNCTION-BY-FUNCTION MANUAL")
    r_badge.font.name = "Segoe UI"
    r_badge.font.size = Pt(9)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(0, 245, 255)

    p_title = c_hero.add_paragraph()
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Every Single Button, Knob & Tool Explained Simply")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = c_hero.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(0)
    r_sub = p_sub.add_run("An exhaustive, easy-to-learn dictionary covering all interface controls, hotkeys, synthesis parameters, and mixing effects.")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(203, 213, 225)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==============================================================
    # SECTION 1: TOP BAR & TRANSPORT
    # ==============================================================
    add_section_header("1. Top Bar & Transport Controls", "The command center that controls song playback, tempo, master volume, and exports.")
    add_function_table([
        ("Play / Pause (Spacebar)", "Starts playback from the current playhead position or pauses it.", "Press Spacebar or click the green Play triangle."),
        ("Stop (Ctrl + Space)", "Halts playback completely and rewinds playhead back to Bar 1.", "Click square Stop icon or press Ctrl + Space."),
        ("Record Arm ('R')", "Primes the track to capture live notes from keyboard or vocals from microphone.", "Click red Record circle or press 'R', then press Spacebar to start take."),
        ("Loop Region Mode ('L')", "Cycles playback continuously over your marked loop bars.", "Click loop icon or press 'L'. Drag along ruler to set loop length."),
        ("Tempo / BPM Scrub", "Controls song speed (Beats Per Minute). 120-140 BPM for Trap/House; 80-95 BPM for Hip-Hop/Lo-Fi.", "Click and drag the BPM number up/down or double click to type exact value."),
        ("Master Volume Slider", "Global volume fader that controls loudness to your speakers/headphones.", "Drag horizontal slider in top bar. Keep below red clipping level!"),
        ("LUFS Loudness Meter", "Measures true broadcast perceived loudness (-14 LUFS is the Spotify standard).", "Watch readout during loudest part of your song. Aim for -14 LUFS."),
        ("Root Key & Scale", "Sets the musical key (e.g. C Minor). Syncs scale highlighting across piano roll.", "Click scale dropdown in top bar; pick your song's key."),
        ("Export WAV Button", "Renders finished song into studio-quality 24-bit / 48 kHz uncompressed stereo WAV.", "Click 'Export' button in top bar or File → Export WAV."),
        ("Export Stems", "Renders separate WAV files for every track (Drums.wav, Bass.wav, Vocal.wav).", "Go to File → Export Stems. Crucial for mixing engineers."),
    ])

    # ==============================================================
    # SECTION 2: STATUS BAR & LIVE TELEMETRY
    # ==============================================================
    add_section_header("2. Status Bar & Hint Panel (Bottom Bar)", "Gives live real-time feedback so you are never lost.")
    add_function_table([
        ("Live Hint Display (💡)", "Shows the exact name, parameter value, and keyboard shortcut of whatever you hover over.", "Just hover your mouse over any knob, slider, or button anywhere in the app!"),
        ("Voice Pool Counter", "Shows how many synthesizer notes are currently sounding at the same instant.", "Monitor polyphony headroom when playing large chords or long synth pads."),
        ("CPU Load Meter", "Real-time DSP audio processing load percentage.", "If it exceeds 85%, increase buffer size in Options → Audio Settings."),
        ("Latency Display (ms)", "Audio round-trip delay time (e.g., 5.3 ms).", "Keep below 10 ms for comfortable vocal and keyboard recording."),
        ("Precision Clock", "Shows current time position in Bar : Beat : 16th : Tick format.", "Use to verify beat grid alignment down to 1/16th note fractions."),
    ])

    # ==============================================================
    # SECTION 3: PRO SAMPLE BROWSER
    # ==============================================================
    add_section_header("3. Pro Sample Browser (Left Panel)", "Your library for auditioning and finding samples, presets, and instruments.")
    add_function_table([
        ("Audition Player (🔊)", "Instantly plays any audio file in real time when clicked before loading.", "Single-click on any Kick, Snare, Hi-Hat, or Melody sample to hear it."),
        ("Add Button (+)", "Loads the auditioned sample into a brand new Channel Rack pad.", "Click the '+' button next to the sound name."),
        ("Packs Tab", "Contains the curated stock sound library (Kicks, Snares, Hats, 808s, Percussion).", "Click 'Packs' tab to browse sample folders."),
        ("Project Tab", "Shows all audio files, recorded takes, and patterns currently used in the project.", "Click 'Project' tab to inspect session memory."),
        ("Favorites Tab (⭐)", "Stores your favorite go-to sounds for 1-click access.", "Click the star icon on any sound to pin it to Favorites."),
        ("Search Filter Bar", "Instant search that filters hundreds of samples by text.", "Type 'snare', '808', or 'drill' into the search bar at the top."),
    ])

    # ==============================================================
    # SECTION 4: CHANNEL RACK PRO (DRUMS & STEP SEQUENCER - F6)
    # ==============================================================
    add_section_header("4. Channel Rack Pro (Press F6)", "The heart of beatmaking: a 16-step matrix with groove swing and graph editor.")
    add_function_table([
        ("16-Step Color Buttons", "Triggers drum hits. 4-step blocks alternate colors (beats 1, 2, 3, 4).", "Left-click on any step button to activate hit (turns highlighted)."),
        ("Mute / Solo LED (●)", "Green circle icon. Silences or isolates that drum pad.", "Left-click to Mute. Right-click or Alt+click to Solo."),
        ("Pan Knob (PAN)", "Positions drum sound across the stereo spectrum (Left / Right).", "Click and drag knob up/down. Double-click to reset to Center."),
        ("Volume Knob (VOL)", "Controls individual volume level for that sample pad.", "Click and drag knob up/down to balance kick vs snare volume."),
        ("Mixer Track (TRK)", "Routes channel output into a designated Mixer insert slot (1 to 10).", "Click number and drag to set target Mixer channel."),
        ("Global Swing Knob", "Delays alternate 16th notes to give drum loops a human, bouncy groove.", "Turn up from 0% to 55%–65% for hip-hop, trap, house, and boom-bap bounce."),
        ("Graph Editor Toggle", "Opens per-step micro-editing bars beneath every step button.", "Click 'Graph Editor' in toolbar; select VEL, PAN, or PIT."),
        ("Graph: Velocity (VEL)", "Adjusts the hit loudness of each individual step.", "Click and drag height of the vertical velocity bar below the step."),
        ("Graph: Pan (PAN)", "Pans individual hits left or right (great for rolling hi-hats).", "Drag bar up (Right) or down (Left) on individual steps."),
        ("Graph: Pitch (PIT)", "Tunes pitch of individual hits up or down.", "Use to create pitch-gliding 808s and tuned tom rolls."),
    ])

    # ==============================================================
    # SECTION 5: PIANO ROLL PRO (MELODIES & CHORDS - F7)
    # ==============================================================
    add_section_header("5. Piano Roll Pro (Press F7)", "Expressive note grid for composition with chord stamper and arpeggiator.")
    add_function_table([
        ("Draw Note", "Places a new musical note on the grid.", "Left-click on any cell in the piano roll grid."),
        ("Resize Note", "Shortens or lengthens the note duration.", "Hover mouse over right edge of note until cursor changes, then drag."),
        ("Delete Note", "Removes note from pattern.", "Right-click directly on the note."),
        ("Scale Highlighting", "Tints rows belonging to your song's scale so you never hit a wrong note.", "Choose key (e.g. A) and scale (e.g. Minor) in toolbar; only play colored rows!"),
        ("Chord Stamper", "Builds full 3-to-4 note chords in 1 click (Major, Minor, 7th, 9th, Sus4).", "Click 'Chord Stamp', pick chord type, click grid to stamp chord!"),
        ("Guitar Strum Tool", "Micro-offsets chord note onsets sequentially from low to high strings.", "Select notes, click 'Strum'. Transforms block chords into natural guitar strums."),
        ("Arpeggiator Engine", "Breaks chords into flowing arpeggiated melodic runs (Up, Down, UpDown, Random).", "Highlight notes, click 'Arp', select direction and 1/16 note rate."),
        ("Humanizer Tool", "Injects random micro-timing offsets and natural velocity drift.", "Select notes, click 'Humanize'. Makes melodies sound hand-played by a human."),
        ("Flam Tool", "Adds rapid double-tap strike articulations to notes.", "Select snare/trap notes, click 'Flam' for fast double hits."),
    ])

    # ==============================================================
    # SECTION 6: EDISON AUDIO SUITE (SAMPLING & SLICING)
    # ==============================================================
    add_section_header("6. Edison Audio Waveform & Slicing Suite", "Visual waveform editor for sample surgery, loop chopping, and normalization.")
    add_function_table([
        ("Load Sample", "Imports any external audio file (.wav, .mp3) into Edison for editing.", "Drag and drop file into Edison or click 'Load Sample' button."),
        ("Zoom In / Out", "Magnifies waveform down to individual transient cycles.", "Scroll mouse wheel over waveform canvas or click Zoom buttons."),
        ("Drag-Selection", "Highlights a specific portion of the waveform for processing.", "Left-click and drag across the waveform canvas."),
        ("Normalize (0 dB)", "Maximizes volume of selected region so highest peak reaches 0 dBFS.", "Click 'Normalize'. Instantly fixes quiet sample recordings!"),
        ("Reverse", "Flips selected audio region backwards in time.", "Select region, click 'Reverse'. Great for reverse cymbal swells!"),
        ("Trim", "Destructively deletes all audio outside your highlighted region.", "Select desired audio, click 'Trim'. Cuts out unwanted dead silence."),
        ("Silence", "Mutes selected region completely to 0 volume.", "Select unwanted cough or background noise, click 'Silence'."),
        ("Fade In / Fade Out", "Applies smooth exponential volume ramps to prevent clicks at sample edges.", "Highlight start or end of loop, click 'Fade In' or 'Fade Out'."),
        ("Transient Auto-Slice", "Analyzes audio energy peaks and places slice markers on every drum hit.", "Click 'Auto Slice'. Chops loops cleanly at every beat automatically!"),
        ("Dump to Rack", "Turns sliced hits into playable trigger pads in the Channel Rack.", "Click 'Dump to Rack'. You can now play the chops on your keyboard!"),
    ])

    # ==============================================================
    # SECTION 7: MIXER PRO & 10-SLOT FX INSPECTOR (F9)
    # ==============================================================
    add_section_header("7. Mixer Pro & 10-Slot FX Inspector (Press F9)", "Channel strips, routing, and studio-grade DSP plugins.")
    add_function_table([
        ("Channel Strip Fader", "Controls channel volume from -∞ dB to +6 dB.", "Drag vertical fader thumb up or down."),
        ("Channel Pan Pot", "Positions channel sound Left / Center / Right.", "Rotate circular pan knob above fader."),
        ("Mute ('M') / Solo ('S')", "Silences channel or isolates it solo during mixdown.", "Click 'M' to mute. Click 'S' to hear only that channel."),
        ("10 FX Insert Slots", "Chain up to 10 serial audio effects on any mixer channel.", "Select a channel strip, look at the right inspector panel, click an empty slot."),
        ("Slot Enable LED", "Green light next to effect slot; bypasses plugin without deleting settings.", "Click green LED to toggle plugin bypass on/off."),
        ("Wet/Dry Mix Slider", "Blends original unprocessed signal with processed effect output.", "Drag horizontal slider next to effect slot (0% = Dry, 100% = Wet)."),
        ("SoftClipper Plugin", "Polynomial saturation transfer curve (analog warmth, punchy drums).", "Insert on Master or Drum bus. Lets 808s slam without digital distortion!"),
        ("StereoShaper Plugin", "Haas micro-delay + Mid/Side stereo width expander.", "Insert on synths, guitars, or pads to make them sound massive in 3D stereo."),
        ("Vintage Chorus Plugin", "Dual modulated delay lines for 80s shimmer and chorus warmth.", "Insert on electric pianos, guitars, or vocal layers."),
        ("Parametric EQ", "Shapes frequency balance (Low-Cut, Mid Scoop, High Shelf).", "Use Low-Cut at 30 Hz to remove muddy rumble from non-bass instruments."),
        ("Compressor Plugin", "Reduces dynamic range; glues drums together.", "Set Threshold so needle catches loudest peaks; add Makeup Gain."),
        ("Reverb & Delay Plugins", "Algorithmic space simulation and ping-pong echoes.", "Insert on vocal/lead tracks; set Wet slider to 15%–25%."),
    ])

    # ==============================================================
    # SECTION 8: PLAYLIST / ARRANGEMENT (F5)
    # ==============================================================
    add_section_header("8. Playlist / Arrangement Timeline (Press F5)", "Where clips, patterns, and vocal takes are organized into complete songs.")
    add_function_table([
        ("Draw Tool ('P')", "Paints pattern blocks or audio clips onto track lanes.", "Select clip from left list, click on timeline grid cells to place."),
        ("Slice / Razor Tool ('C')", "Splits any placed clip into two separate pieces at the cursor line.", "Select Razor tool, click directly on clip where you want to cut."),
        ("Select Tool ('E')", "Box-selects multiple clips across multiple tracks.", "Click and drag selection rectangle over clips to move them together."),
        ("Duplicate ('Ctrl + B')", "Duplicates selected clips and pastes them flush to the next bar.", "Select clips and press Ctrl + B to quickly build out verses and choruses."),
        ("Section Markers", "Labels song sections along the timeline ruler (Intro, Verse, Chorus, Outro).", "Click 'Add Marker' on timeline ruler. Click marker to jump playhead!"),
        ("Track Mute / Solo", "Silences or isolates an entire arrangement track lane.", "Click M/S icons on the track header on the left side of timeline."),
    ])

    # ==============================================================
    # SECTION 9: VOCAL & AUDIO RECORDING STUDIO
    # ==============================================================
    add_section_header("9. Vocal & Live Instrument Recording Studio", "Captures vocals and live instruments with zero timing delay.")
    add_function_table([
        ("Input Device Selector", "Chooses which connected microphone or audio interface to capture.", "Click input dropdown in Vocal workspace; pick your mic."),
        ("Live Monitor Toggle", "Routes microphone audio into your headphones in real time.", "Click 'Monitor' button. (Use headphones to avoid speaker feedback!)"),
        ("Record Arm Button", "Prepares track to capture audio.", "Click 'Record Arm' or press 'R'. Button turns red."),
        ("Take Management List", "Stores every recorded vocal attempt as an independent take.", "View all takes in Take List on the left side of Vocal workspace."),
        ("Automatic Latency Alignment", "Compensates for hardware audio delay, locking voice onto the beat grid.", "Handled automatically by Cadence DSP engine in the background!"),
        ("Take Comping", "Auditioning takes and assembling the best sections into a master take.", "Click take names in list to compare performances and select the best one."),
    ])

    # ==============================================================
    # SECTION 10: SYNTH LAB (SUBTRACTIVE SYNTHESIS)
    # ==============================================================
    add_section_header("10. Synth Lab (Subtractive Sound Design)", "Built-in polyphonic synthesizer for basses, plucks, leads, and pads.")
    add_function_table([
        ("Oscillator Waveform (OSC)", "Chooses raw waveform shape: Saw (bright), Square (hollow), Sine (pure), Tri (soft).", "Click OSC waveform icon to select base sound character."),
        ("Filter Cutoff Knob (CUT)", "Controls low-pass filter frequency; removes highs for dark or warm sounds.", "Rotate Cutoff knob down for warm sub-bass; up for bright leads."),
        ("Resonance Knob (RES)", "Boosts frequencies right at the cutoff point for synth squelch / bite.", "Turn up slightly to make synth filter sweeps cut through the mix."),
        ("Attack (A)", "Time taken for note to reach full volume after key press.", "0 ms for fast punchy plucks/bass; 300-800 ms for slow ambient pads."),
        ("Decay (D)", "Time taken to drop from peak volume down to sustain level.", "Set to 100-300 ms for percussive pluck sounds."),
        ("Sustain (S)", "Volume level held while key remains pressed down.", "100% for organs and leads; 0% for bells and plucks."),
        ("Release (R)", "Time taken for note to fade to silence after key is released.", "50 ms for tight bass; 1000 ms for long atmospheric washes."),
    ])

    add_callout(
        "Save this manual as a quick reference on your desktop! Whenever you are unsure of what a button or knob does in Cadence, simply hover over it in the DAW to read the live Hint Bar, or search this document.",
        title="MASTER TIP",
        icon="🚀"
    )

    out_path = os.path.abspath("Cadence_DAW_Complete_Function_Reference.docx")
    doc.save(out_path)
    print(f"SUCCESS: Saved {out_path}")

if __name__ == "__main__":
    build_complete_reference()
